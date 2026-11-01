"""Importación del «Reporte Prestador» del portal ISL (QA-004, QA UTalca 22-10).

El CEPA descarga cada quincena desde el portal del ISL un Excel con la vigencia de las
órdenes de atención (ODA/OPA/OPAEP/Ambulatoria). Este servicio lo lee y, por cada fila:

1. Busca el ingreso: primero por N° Siniestro contra ``ingreso.numero_siniestro``; si no
   hay, por RUT del trabajador contra ``paciente.rut`` (respaldo: en el archivo real la
   columna «Rut Trabajador» viene casi siempre vacía o con ``sin_id_<n>``).
   Si hay varios ingresos candidatos se prefiere el activo y, entre iguales, el más reciente.
2. Hace upsert de la ODA por (ingreso, identificador = N° de Orden): subir la misma planilla
   dos veces no duplica nada.
3. Recalcula la vigente de cada ingreso tocado: una sola por ingreso, la de vencimiento más
   tardío (desempate: la que el ISL reporta VIGENTE, luego la más nueva). Es la misma regla
   que ``registrar_oda`` y la que leen el dashboard y el reporte de ODAS vencidas
   (vigente + fecha pasada = vencida).

La alerta ODA_POR_VENCER solo considera ingresos activos (ver ``alertas._construir_hitos_oda``).
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select, true
from sqlalchemy.orm import Session

from app.domain.enums import EstadoCaso
from app.models.ingreso import Ingreso
from app.models.oda import Oda
from app.models.paciente import Paciente
from app.util.rut import RutInvalidoError, normalizar_rut

HOJA_ESPERADA = "Reporte Prestador"
MAX_FILAS_ENCABEZADO = 15  # filas en que se busca el encabezado (por si el ISL agrega títulos)
LOTE_IN = 500  # Oracle limita las listas IN a 1000 elementos
# Topes contra planillas desmedidas o «zip bombs»: el límite de 10 MB del endpoint es sobre
# el archivo comprimido; un xlsx real del ISL (~2.600 filas) pesa ~1 MB descomprimido.
MAX_DESCOMPRIMIDO_BYTES = 100 * 1024 * 1024
MAX_FILAS = 50_000

# Clave normalizada del encabezado → nombre lógico de la columna.
_COLUMNAS = {
    "ndeorden": "numero_orden",
    "nroorden": "numero_orden",
    "numerodeorden": "numero_orden",
    "tipodeorden": "tipo_orden",
    "iniciovigencia": "inicio",
    "finvigencia": "fin",
    "fechacese": "cese",
    "estadoorden": "estado",
    "nsiniestro": "siniestro",
    "nrosiniestro": "siniestro",
    "numerodesiniestro": "siniestro",
    "ruttrabajador": "rut_trabajador",
}
_OBLIGATORIAS = ("numero_orden", "inicio", "fin", "siniestro")
_NOMBRES_VISIBLES = {
    "numero_orden": "N° de Orden",
    "inicio": "Inicio Vigencia",
    "fin": "Fin Vigencia",
    "siniestro": "N° Siniestro",
}


class ArchivoInvalidoError(ValueError):
    """El archivo no es un xlsx legible o no trae las columnas del Reporte Prestador."""


@dataclass
class FilaReporte:
    fila: int  # número de fila en la planilla (1 = primera fila de Excel)
    numero_orden: str
    tipo_orden: str | None
    inicio: date | None
    fin: date
    cese: date | None
    estado: str | None
    siniestro: str | None
    rut_trabajador: str | None

    @property
    def fecha_vencimiento(self) -> date:
        """Fin Vigencia, o Fecha Cese si existe y es anterior."""
        if self.cese is not None and self.cese < self.fin:
            return self.cese
        return self.fin

    @property
    def vigente_isl(self) -> bool:
        return (self.estado or "").upper() == "VIGENTE"


@dataclass
class ErrorFila:
    fila: int
    motivo: str


@dataclass
class ResultadoParseo:
    filas: list[FilaReporte] = field(default_factory=list)
    errores: list[ErrorFila] = field(default_factory=list)
    filas_leidas: int = 0


# ── Parser ─────────────────────────────────────────────────────────────────────


def _clave(texto: object) -> str:
    """Normaliza un encabezado: sin tildes, sin símbolos, minúsculas ('N° de Orden' → 'ndeorden')."""
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _texto(valor: object) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    s = str(valor).strip()
    return s or None


def _fecha(valor: object) -> date | None:
    """Acepta date/datetime de Excel o texto dd-mm-aaaa / dd/mm/aaaa / aaaa-mm-dd."""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    if not s:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha inválida «{s}»")


def parsear_reporte(contenido: bytes) -> ResultadoParseo:
    """Lee el xlsx y devuelve las filas válidas y los errores por fila. No toca la BD."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as zf:
            descomprimido = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise ArchivoInvalidoError("El archivo no es un Excel (.xlsx) válido.") from exc
    if descomprimido > MAX_DESCOMPRIMIDO_BYTES:
        raise ArchivoInvalidoError("El archivo es demasiado grande para un Reporte Prestador.")

    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl lanza tipos variados ante archivos rotos
        raise ArchivoInvalidoError("El archivo no es un Excel (.xlsx) válido.") from exc
    try:
        ws = wb[HOJA_ESPERADA] if HOJA_ESPERADA in wb.sheetnames else wb.worksheets[0]
        filas = ws.iter_rows(values_only=True)

        indices: dict[str, int] | None = None
        num_fila = 0
        for num_fila, fila in enumerate(filas, start=1):
            mapa = {}
            for i, celda in enumerate(fila):
                nombre = _COLUMNAS.get(_clave(celda))
                if nombre and nombre not in mapa:
                    mapa[nombre] = i
            if "numero_orden" in mapa:
                indices = mapa
                break
            if num_fila >= MAX_FILAS_ENCABEZADO:
                break
        if indices is None:
            raise ArchivoInvalidoError(
                "No se encontró el encabezado del Reporte Prestador (columna «N° de Orden»)."
            )
        faltan = [_NOMBRES_VISIBLES[c] for c in _OBLIGATORIAS if c not in indices]
        if faltan:
            raise ArchivoInvalidoError("Faltan columnas en el archivo: " + ", ".join(faltan) + ".")

        resultado = ResultadoParseo()
        for num_fila, fila in enumerate(filas, start=num_fila + 1):
            if fila is None or all(_texto(c) is None for c in fila):
                continue  # filas vacías (p. ej. al final de la hoja)
            resultado.filas_leidas += 1
            if resultado.filas_leidas > MAX_FILAS:
                raise ArchivoInvalidoError(
                    f"El archivo supera el máximo de {MAX_FILAS:,} filas.".replace(",", ".")
                )

            def celda(nombre: str, _fila=fila) -> object:
                i = indices.get(nombre)
                return _fila[i] if i is not None and i < len(_fila) else None

            numero = _texto(celda("numero_orden"))
            if numero is None:
                resultado.errores.append(ErrorFila(num_fila, "Sin N° de Orden."))
                continue
            if len(numero) > 60:
                resultado.errores.append(ErrorFila(num_fila, "N° de Orden demasiado largo."))
                continue
            try:
                inicio = _fecha(celda("inicio"))
                fin = _fecha(celda("fin"))
                cese = _fecha(celda("cese"))
            except ValueError as exc:
                resultado.errores.append(ErrorFila(num_fila, f"Orden {numero}: {exc}."))
                continue
            if fin is None:
                resultado.errores.append(ErrorFila(num_fila, f"Orden {numero}: sin Fin Vigencia."))
                continue
            siniestro = _texto(celda("siniestro"))
            rut = _texto(celda("rut_trabajador"))
            if siniestro is None and rut is None:
                resultado.errores.append(
                    ErrorFila(num_fila, f"Orden {numero}: sin N° Siniestro ni RUT del trabajador.")
                )
                continue
            resultado.filas.append(
                FilaReporte(
                    fila=num_fila,
                    numero_orden=numero,
                    tipo_orden=_texto(celda("tipo_orden")),
                    inicio=inicio,
                    fin=fin,
                    cese=cese,
                    estado=(_texto(celda("estado")) or "").upper() or None,
                    siniestro=siniestro,
                    rut_trabajador=rut,
                )
            )
        return resultado
    finally:
        wb.close()


# ── Match con ingresos ─────────────────────────────────────────────────────────


def _norm_siniestro(s: str | None) -> str | None:
    if s is None:
        return None
    s = s.strip().upper()
    return (s.lstrip("0") or "0") if s.isdigit() else (s or None)


def _rut_canonico(s: str | None) -> str | None:
    if not s:
        return None
    try:
        return normalizar_rut(s)
    except RutInvalidoError:
        return None  # p. ej. «sin_id_1141409» o el RUT del prestador mal digitado


def _mejor(ingresos: list[Ingreso]) -> Ingreso:
    """Entre varios ingresos candidatos: el activo y, entre iguales, el más reciente."""
    return max(
        ingresos,
        key=lambda i: (i.estado == EstadoCaso.ACTIVO.value, i.fecha_ingreso, i.id),
    )


def _lotes(valores: list, n: int = LOTE_IN):
    for k in range(0, len(valores), n):
        yield valores[k : k + n]


def _indices_ingresos(
    db: Session, filas: list[FilaReporte]
) -> tuple[dict[str, list[Ingreso]], dict[str, list[Ingreso]]]:
    por_siniestro: dict[str, list[Ingreso]] = {}
    siniestros = {_norm_siniestro(f.siniestro) for f in filas} - {None}
    if siniestros:
        # numero_siniestro puede venir con ceros a la izquierda o espacios: se normaliza en
        # Python. Los ingresos con siniestro son a lo más unos miles: una sola consulta.
        for ing in db.scalars(select(Ingreso).where(Ingreso.numero_siniestro.isnot(None))):
            clave = _norm_siniestro(ing.numero_siniestro)
            if clave in siniestros:
                por_siniestro.setdefault(clave, []).append(ing)

    por_rut: dict[str, list[Ingreso]] = {}
    ruts = sorted({_rut_canonico(f.rut_trabajador) for f in filas} - {None})
    for lote in _lotes(ruts):
        consulta = (
            select(Paciente.rut, Ingreso)
            .join(Ingreso, Ingreso.paciente_id == Paciente.id)
            .where(Paciente.rut.in_(lote))
        )
        for rut, ing in db.execute(consulta):
            por_rut.setdefault(rut, []).append(ing)
    return por_siniestro, por_rut


# ── Importación ────────────────────────────────────────────────────────────────


@dataclass
class FilaSinPaciente:
    fila: int
    numero_orden: str
    tipo_orden: str | None
    numero_siniestro: str | None
    rut_trabajador: str | None


@dataclass
class OdaPorVencer:
    oda_id: int
    ingreso_id: int
    folio: str
    numero_orden: str
    fecha_vencimiento: date


@dataclass
class ResumenImportacion:
    filas_leidas: int = 0
    creadas: int = 0
    actualizadas: int = 0
    sin_cambios: int = 0
    match_por_siniestro: int = 0
    match_por_rut: int = 0
    sin_paciente: list[FilaSinPaciente] = field(default_factory=list)
    errores: list[ErrorFila] = field(default_factory=list)
    por_vencer_paciente_activo: list[OdaPorVencer] = field(default_factory=list)
    ventana_dias: int = 0
    # (oda, acción) para auditar; no se expone en la respuesta.
    cambios: list[tuple[Oda, str]] = field(default_factory=list)


def _ventana_oda(db: Session) -> int:
    from app.domain.enums_alertas import TipoAlerta
    from app.services.config_alertas import cargar_ventanas

    return int(cargar_ventanas(db)[TipoAlerta.ODA_POR_VENCER.value]["dias"])


def importar_reporte(
    db: Session, contenido: bytes, *, hoy: date | None = None
) -> ResumenImportacion:
    """Aplica el Reporte Prestador sobre la BD (solo flush; el caller decide commit/rollback).

    Lanza ArchivoInvalidoError si el archivo no se puede leer.
    """
    hoy = hoy or date.today()
    parseo = parsear_reporte(contenido)
    resumen = ResumenImportacion(filas_leidas=parseo.filas_leidas, errores=list(parseo.errores))
    por_siniestro, por_rut = _indices_ingresos(db, parseo.filas)

    # Filas con ingreso resuelto, agrupadas por ingreso.
    por_ingreso: dict[int, list[FilaReporte]] = {}
    ingresos: dict[int, Ingreso] = {}
    for f in parseo.filas:
        candidatos = por_siniestro.get(_norm_siniestro(f.siniestro) or "")
        if candidatos:
            resumen.match_por_siniestro += 1
        else:
            candidatos = por_rut.get(_rut_canonico(f.rut_trabajador) or "")
            if candidatos:
                resumen.match_por_rut += 1
        if not candidatos:
            resumen.sin_paciente.append(
                FilaSinPaciente(f.fila, f.numero_orden, f.tipo_orden, f.siniestro, f.rut_trabajador)
            )
            continue
        ing = _mejor(candidatos)
        ingresos[ing.id] = ing
        por_ingreso.setdefault(ing.id, []).append(f)

    for ingreso_id, filas in por_ingreso.items():
        existentes = {
            o.identificador: o
            for o in db.scalars(select(Oda).where(Oda.ingreso_id == ingreso_id))
        }
        vigente_isl: set[str] = set()
        for f in filas:
            if f.vigente_isl:
                vigente_isl.add(f.numero_orden)
            oda = existentes.get(f.numero_orden)
            if oda is None:
                oda = Oda(
                    ingreso_id=ingreso_id,
                    identificador=f.numero_orden,
                    fecha_registro=f.inicio,
                    fecha_vencimiento=f.fecha_vencimiento,
                    vigente=False,
                )
                db.add(oda)
                existentes[f.numero_orden] = oda
                resumen.creadas += 1
                resumen.cambios.append((oda, "CREATE"))
            elif (oda.fecha_registro, oda.fecha_vencimiento) != (f.inicio, f.fecha_vencimiento):
                oda.fecha_registro = f.inicio
                oda.fecha_vencimiento = f.fecha_vencimiento
                resumen.actualizadas += 1
                resumen.cambios.append((oda, "UPDATE"))
            else:
                resumen.sin_cambios += 1
        db.flush()

        # Una vigente por ingreso: la de vencimiento más tardío.
        elegida = max(
            existentes.values(),
            key=lambda o: (o.fecha_vencimiento, o.identificador in vigente_isl, o.id),
        )
        for o in existentes.values():
            o.vigente = o is elegida
    db.flush()

    resumen.ventana_dias = _ventana_oda(db)
    limite = hoy + timedelta(days=resumen.ventana_dias)
    for ingreso_id in por_ingreso:
        ing = ingresos[ingreso_id]
        if ing.estado != EstadoCaso.ACTIVO.value:
            continue
        for o in db.scalars(
            select(Oda).where(Oda.ingreso_id == ingreso_id, Oda.vigente == true())
        ):
            if hoy <= o.fecha_vencimiento <= limite:
                resumen.por_vencer_paciente_activo.append(
                    OdaPorVencer(o.id, ingreso_id, ing.folio, o.identificador, o.fecha_vencimiento)
                )
    resumen.por_vencer_paciente_activo.sort(key=lambda p: (p.fecha_vencimiento, p.folio))
    return resumen
